from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from rest_framework import status

from clientes.models import Cliente
from pruebas.base import TIPO_EXCEL, PruebaConSesion, PruebaSinSesion
from gastos.models import Gasto
from pedidos.models import Cobro, Pedido, ProductoDelPedido
from productos.models import Producto

URL = '/api/finanzas/'


class FinanzasDelPeriodo(PruebaConSesion):
    """Ingresos, gastos y resultado de un período (CU61, CU62)."""

    def test_desde_posterior_a_hasta_da_400(self):
        """CU61 y CU62: un período al revés responde 400 con mensaje."""
        respuesta = self.client.get(URL, {'desde': '2026-02-01', 'hasta': '2026-01-01'})
        self.assertEqual(respuesta.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(respuesta.data['detail'],
                         'La fecha desde no puede ser posterior a la fecha hasta.')

    def test_los_totales_del_periodo_son_correctos(self):
        """CU61 y CU62: un cobro de $4.000 y un gasto de $1.500 en marzo dan resultado $2.500;
        lo de febrero y abril queda afuera."""
        cliente = Cliente.objects.create(instagram='sofi.delgado', nombre='Sofía')
        producto = Producto.objects.create(
            nombre='Aros Hedwig', precio_actual='5000.00',
            dificultad=Producto.Dificultad.MEDIA,
        )
        pedido = Pedido.objects.create(cliente=cliente)
        ProductoDelPedido.objects.create(pedido=pedido, producto=producto,
                                         precio='5000.00', cantidad=2)
        Cobro.objects.create(pedido=pedido, monto='4000.00', fecha='2026-03-10')
        Cobro.objects.create(pedido=pedido, monto='1000.00', fecha='2026-04-01')   # afuera
        Gasto.objects.create(tipo=Gasto.Tipo.OTRO, monto='1500.00', fecha='2026-03-15')
        Gasto.objects.create(tipo=Gasto.Tipo.OTRO, monto='700.00', fecha='2026-02-28')  # afuera

        respuesta = self.client.get(URL, {'desde': '2026-03-01', 'hasta': '2026-03-31'})
        self.assertEqual(respuesta.status_code, status.HTTP_200_OK)
        self.assertEqual(respuesta.data['totales'], {
            'ingresos': '4000.00',
            'gastos': '1500.00',
            'resultado': '2500.00',
            'cantidad_cobros': 1,
            'cantidad_gastos': 1,
        })
        self.assertEqual(len(respuesta.data['cobros']), 1)
        self.assertEqual(len(respuesta.data['gastos']), 1)


EXPORTAR = f'{URL}exportar/'
MARZO = {'desde': '2026-03-01', 'hasta': '2026-03-31'}


class ExportarSinSesion(PruebaSinSesion):
    """Exportar es un endpoint de gestión (corrección E10)."""

    def test_exportar_sin_sesion_da_403(self):
        """CU61, CU62 y E10: GET /api/finanzas/exportar/ sin sesión responde 403, en PDF y en Excel."""
        for formato in ('pdf', 'excel'):
            respuesta = self.client.get(EXPORTAR, {**MARZO, 'formato': formato})
            self.assertEqual(respuesta.status_code, status.HTTP_403_FORBIDDEN)


class ExportarFinanzas(PruebaConSesion):
    """El período de Finanzas se exporta a PDF y a Excel (CU61, CU62, corrección E10).

    Dos cobros en marzo, uno en efectivo y otro por transferencia, y dos
    gastos en marzo; un cobro y un gasto fuera del período.
    """

    def setUp(self):
        super().setUp()
        cliente = Cliente.objects.create(instagram='sofi.delgado', nombre='Sofía')
        producto = Producto.objects.create(
            nombre='Aros Hedwig', precio_actual='5000.00', dificultad=Producto.Dificultad.MEDIA,
        )
        pedido = Pedido.objects.create(cliente=cliente)
        ProductoDelPedido.objects.create(pedido=pedido, producto=producto, precio='5000.00', cantidad=2)
        Cobro.objects.create(pedido=pedido, monto='4000.00', fecha='2026-03-10', medio=Cobro.Medio.EFECTIVO)
        Cobro.objects.create(pedido=pedido, monto='2500.00', fecha='2026-03-20', medio=Cobro.Medio.TRANSFERENCIA,
                             tipo=Cobro.Tipo.PAGO_RESTANTE)
        Cobro.objects.create(pedido=pedido, monto='1000.00', fecha='2026-04-01')   # afuera
        Gasto.objects.create(tipo=Gasto.Tipo.OTRO, monto='1500.00', fecha='2026-03-15')
        Gasto.objects.create(tipo=Gasto.Tipo.PUBLICIDAD, monto='800.00', fecha='2026-03-02')
        Gasto.objects.create(tipo=Gasto.Tipo.OTRO, monto='700.00', fecha='2026-02-28')  # afuera

    def libro(self, parametros):
        respuesta = self.client.get(EXPORTAR, {**parametros, 'formato': 'excel'})
        self.assertEqual(respuesta.status_code, status.HTTP_200_OK)
        return load_workbook(BytesIO(b''.join(respuesta.streaming_content)))

    def test_sin_fechas_o_con_formato_invalido_da_400(self):
        """CU61, CU62 y E10: sin desde, o con un formato que no es pdf ni excel, responde 400 con mensaje; bien, 200."""
        respuesta = self.client.get(EXPORTAR, {'hasta': '2026-03-31', 'formato': 'pdf'})
        self.assertEqual(respuesta.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(respuesta.data['detail'], 'Falta la fecha desde.')

        respuesta = self.client.get(EXPORTAR, {**MARZO, 'formato': 'word'})
        self.assertEqual(respuesta.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(respuesta.data['detail'], 'El formato tiene que ser pdf o excel.')

        respuesta = self.client.get(EXPORTAR, {**MARZO, 'formato': 'pdf'})
        self.assertEqual(respuesta.status_code, status.HTTP_200_OK)

    def test_el_pdf_se_muestra_y_el_excel_se_descarga(self):
        """CU61, CU62 y E10: el PDF va inline; el Excel, como adjunto, con las hojas Resumen, Cobros y Gastos."""
        respuesta = self.client.get(EXPORTAR, {**MARZO, 'formato': 'pdf'})
        self.assertEqual(respuesta['Content-Type'], 'application/pdf')
        self.assertTrue(respuesta['Content-Disposition'].startswith('inline'))
        self.assertIn('finanzas-2026-03-01-a-2026-03-31.pdf', respuesta['Content-Disposition'])

        respuesta = self.client.get(EXPORTAR, {**MARZO, 'formato': 'excel'})
        self.assertEqual(respuesta['Content-Type'], TIPO_EXCEL)
        self.assertTrue(respuesta['Content-Disposition'].startswith('attachment'))
        libro = load_workbook(BytesIO(b''.join(respuesta.streaming_content)))
        self.assertEqual(libro.sheetnames, ['Resumen', 'Cobros', 'Gastos'])

    def test_los_totales_del_excel_son_los_de_la_pantalla(self):
        """CU61, CU62 y E10: ingresos, gastos y resultado de la hoja Resumen son los de GET /api/finanzas/, y Cobros y Gastos tienen una fila por movimiento."""
        pantalla = self.client.get(URL, MARZO).data['totales']
        libro = self.libro(MARZO)

        # En Resumen, cada total es una fila con el concepto en A y el importe en B.
        importes = {fila[0]: fila[1] for fila in libro['Resumen'].iter_rows(values_only=True) if fila[0]}
        for concepto, clave in (('Ingresos', 'ingresos'), ('Gastos', 'gastos'), ('Resultado', 'resultado')):
            self.assertEqual(Decimal(str(importes[concepto])), Decimal(pantalla[clave]))
        self.assertEqual(Decimal(pantalla['resultado']), Decimal('4200.00'))

        # En Cobros y Gastos, los títulos de las columnas van en la fila 5 y
        # cada movimiento en una fila desde la 6.
        self.assertEqual(libro['Cobros'].max_row - 5, pantalla['cantidad_cobros'])
        self.assertEqual(libro['Gastos'].max_row - 5, pantalla['cantidad_gastos'])
        self.assertEqual(pantalla['cantidad_cobros'], 2)
