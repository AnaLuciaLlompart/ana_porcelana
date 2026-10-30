"""Exportar un listado a PDF o a Excel (corrección E10 del tutor).

Se exporta lo mismo que muestra el listado: los registros salen del
get_queryset() del ViewSet, con sus filtros, su búsqueda y su orden, pero
sin paginar. Cada ViewSet que hereda ExportarMixin declara solo qué
columnas lleva y cómo se describen sus filtros; el armado de los dos
archivos está escrito una vez, acá.

El armado está en piezas (el encabezado, una tabla, el documento con su
pie; en Excel, el encabezado y una tabla dentro de una hoja) para que un
archivo con más de una tabla, como el de Finanzas, use las mismas.
"""

from collections import namedtuple
from io import BytesIO
from xml.sax.saxutils import escape

from django.http import FileResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import simpleSplit
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from rest_framework.decorators import action
from rest_framework.exceptions import ParseError

from pedidos.presentacion import plata


# Una columna del archivo. valor es una función que recibe un registro y
# devuelve lo que va en la celda; tipo decide cómo se escribe (texto,
# plata, fecha, mes o numero), y ancho es su parte del ancho de la tabla
# del PDF.
Columna = namedtuple('Columna', ['titulo', 'valor', 'tipo', 'ancho'])

# Para el tipo 'mes', que se escribe "Octubre 2026".
MESES_LARGOS = (
    'Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio',
    'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre',
)


# ---------------------------------------------------------------------
# Lo que comparten los dos formatos
# ---------------------------------------------------------------------

def escribir_fecha(valor):
    """Una fecha como dd/mm/aaaa, o vacío si no hay."""
    return valor.strftime('%d/%m/%Y') if valor else ''


def describir_rango(nombre, desde, hasta, escribir=plata):
    """Un filtro de rango en palabras, para describir_filtros().

    'Total: $5.000 – $90.000', 'Total: desde $5.000' o 'Total: hasta
    $90.000', como el chip de la pantalla; None si no hay ningún extremo.
    escribir dice cómo se escribe cada extremo: plata para los importes,
    escribir_fecha para un período.
    """
    if desde is not None and hasta is not None:
        return f'{nombre}: {escribir(desde)} – {escribir(hasta)}'
    if desde is not None:
        return f'{nombre}: desde {escribir(desde)}'
    if hasta is not None:
        return f'{nombre}: hasta {escribir(hasta)}'
    return None


def _como_texto(valor, tipo):
    """El valor de una celda escrito para leer, como va en el PDF."""
    if valor is None or valor == '':
        return ''
    if tipo == 'plata':
        # plata() escribiría "$-484.970"; el signo va antes del peso.
        return '-' + plata(-valor) if valor < 0 else plata(valor)
    if tipo == 'fecha':
        return escribir_fecha(valor)
    if tipo == 'mes':
        return f'{MESES_LARGOS[valor.month - 1]} {valor.year}'
    return str(valor)


# ---------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------
# Con la paleta y la letra del comprobante (pedidos/comprobante.py). Las
# medidas van en puntos de PDF y no pasan por px(): el comprobante copia
# un diseño en píxeles, y estos listados no tienen uno.

ROSA_PROFUNDO = HexColor('#8C5A66')
ROSA_CLARO = HexColor('#F0E2E4')
TEXTO = HexColor('#3D3238')
TEXTO_SECUNDARIO = HexColor('#857078')
BORDE = HexColor('#EBE0E2')

# Una A4 apaisada, porque los listados son anchos, y lo que queda entre
# los dos márgenes de los costados.
HOJA = landscape(A4)
MARGEN = 32
ANCHO_UTIL = HOJA[0] - 2 * MARGEN

# SimpleDocTemplate deja 6 puntos de relleno por dentro de cada margen, y
# los textos arrancaban 6 puntos más adentro que las tablas. Se descuentan
# de los márgenes de los costados, como en el comprobante, así textos,
# tablas y pie quedan alineados en MARGEN.
RELLENO_DEL_MARCO = 6

MARCA = ParagraphStyle('marca', fontName='Helvetica-Bold', fontSize=20, leading=24, textColor=ROSA_PROFUNDO)
TITULO = ParagraphStyle('titulo', fontName='Helvetica-Bold', fontSize=15, leading=19, textColor=TEXTO)
DETALLE = ParagraphStyle('detalle', fontName='Helvetica', fontSize=10, leading=14, textColor=TEXTO_SECUNDARIO)
COLUMNA = ParagraphStyle('columna', fontName='Helvetica-Bold', fontSize=8, leading=10, textColor=ROSA_PROFUNDO)
COLUMNA_DERECHA = ParagraphStyle('columna_derecha', parent=COLUMNA, alignment=TA_RIGHT)
CELDA = ParagraphStyle('celda', fontName='Helvetica', fontSize=9, leading=12, textColor=TEXTO)
CELDA_DERECHA = ParagraphStyle('celda_derecha', parent=CELDA, alignment=TA_RIGHT)


def _parrafo(texto, estilo):
    """Un párrafo con el texto protegido, como en el comprobante: un & o un <
    sueltos en un nombre romperían el PDF."""
    return Paragraph(escape(texto), estilo)


# El pie de cada hoja. BASE_PIE es la altura del renglón de abajo, medida
# desde el borde de la hoja; los filtros pueden ocupar hasta tres renglones
# hacia arriba, y encima va la raya. El margen de abajo del documento deja
# ese lugar libre, para que la tabla no se escriba encima.
BASE_PIE = 16
RENGLON_PIE = 10
RENGLONES_PIE = 3
ALTURA_RAYA = BASE_PIE + RENGLONES_PIE * RENGLON_PIE + 4
MARGEN_INFERIOR = ALTURA_RAYA + 12


def _pie(texto_izquierda, fecha):
    """Arma la función que dibuja el pie, que ReportLab llama en cada hoja.

    A la izquierda va texto_izquierda (los filtros de un listado, el
    período de Finanzas), la fecha en el medio y el número de página a la
    derecha. El texto de la izquierda puede ser largo: simpleSplit lo corta
    en renglones que entran en el 40 % del ancho, y se dibujan de abajo
    hacia arriba, así el último queda a la altura de la fecha.
    """
    def dibujar(canvas, documento):
        ancho_hoja = documento.pagesize[0]
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.setFillColor(TEXTO_SECUNDARIO)

        renglones = simpleSplit(texto_izquierda, 'Helvetica', 8, (ancho_hoja - 2 * MARGEN) * 0.4)
        if len(renglones) > RENGLONES_PIE:
            renglones = renglones[:RENGLONES_PIE]
            renglones[-1] += '…'
        for numero, renglon in enumerate(reversed(renglones)):
            canvas.drawString(MARGEN, BASE_PIE + numero * RENGLON_PIE, renglon)

        canvas.drawCentredString(ancho_hoja / 2, BASE_PIE, fecha)
        canvas.drawRightString(ancho_hoja - MARGEN, BASE_PIE, f'Página {documento.page}')

        canvas.setStrokeColor(BORDE)
        canvas.setLineWidth(0.75)
        canvas.line(MARGEN, ALTURA_RAYA, ancho_hoja - MARGEN, ALTURA_RAYA)
        canvas.restoreState()

    return dibujar


def encabezado_pdf(titulo, renglones):
    """Los bloques de arriba de la primera hoja: la marca, el título y los
    renglones de detalle (la fecha, la cantidad, los filtros o el período)."""
    bloques = [_parrafo('Ana Porcelana', MARCA), Spacer(1, 4), _parrafo(titulo, TITULO), Spacer(1, 4)]
    bloques += [_parrafo(renglon, DETALLE) for renglon in renglones]
    bloques.append(Spacer(1, 12))
    return bloques


def tabla_pdf(columnas, filas, ancho=ANCHO_UTIL):
    """Una tabla con los títulos de las columnas en rosa claro y una raya
    debajo de cada fila. Si no entra en la hoja, sigue en la siguiente y
    repite la fila de los títulos."""
    a_la_derecha = [columna.tipo in ('plata', 'numero') for columna in columnas]

    tabla_filas = [[
        _parrafo(columna.titulo.upper(), COLUMNA_DERECHA if derecha else COLUMNA)
        for columna, derecha in zip(columnas, a_la_derecha)
    ]]
    for fila in filas:
        tabla_filas.append([
            _parrafo(_como_texto(valor, columna.tipo), CELDA_DERECHA if derecha else CELDA)
            for valor, columna, derecha in zip(fila, columnas, a_la_derecha)
        ])

    total_anchos = sum(columna.ancho for columna in columnas)
    tabla = Table(
        tabla_filas,
        colWidths=[ancho * columna.ancho / total_anchos for columna in columnas],
        repeatRows=1,
        # Una tabla más angosta que la hoja va contra el margen izquierdo,
        # como los títulos, y no centrada, que es lo que hace ReportLab.
        hAlign='LEFT',
    )
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), ROSA_CLARO),
        ('LINEBELOW', (0, 1), (-1, -1), 0.75, BORDE),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
    ]))
    return tabla


def construir_pdf(titulo, bloques, texto_pie, fecha):
    """Arma el documento con esos bloques y el pie en cada hoja, y lo
    devuelve en memoria, listo para la respuesta."""
    archivo = BytesIO()
    documento = SimpleDocTemplate(
        archivo,
        pagesize=HOJA,
        leftMargin=MARGEN - RELLENO_DEL_MARCO,
        rightMargin=MARGEN - RELLENO_DEL_MARCO,
        topMargin=MARGEN,
        bottomMargin=MARGEN_INFERIOR,
        title=titulo,
        author='Ana Porcelana',
    )
    pie = _pie(texto_pie, fecha)
    documento.build(bloques, onFirstPage=pie, onLaterPages=pie)
    archivo.seek(0)
    return archivo


def armar_pdf(titulo, fecha, cantidad, texto_filtros, columnas, filas):
    """El PDF de un listado: encabezado, la tabla y el pie con los filtros."""
    bloques = encabezado_pdf(titulo, [f'{fecha} · {cantidad}', texto_filtros])
    bloques.append(tabla_pdf(columnas, filas))
    return construir_pdf(titulo, bloques, texto_filtros, fecha)


# ---------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------
# Cada fila se escribe por su número. No se pregunta cuál es la última
# fila usada (max_row): un renglón vacío no cuenta como usado, y la cuenta
# daba una fila menos, así que los títulos de las columnas quedaban fuera
# de la fila pintada y fijada.

# Cuántos caracteres de ancho le toca a cada parte de "ancho" en la hoja.
CARACTERES_POR_ANCHO = 1.6

FORMATO = {
    'plata': '"$"#,##0',
    'fecha': 'DD/MM/YYYY',
    'mes': 'MMMM YYYY',
    'numero': '0',
}


def escribir_encabezado_excel(hoja, titulo, renglones):
    """El título en la fila 1 y cada renglón de detalle debajo, uno por fila.

    Deja una fila vacía, de aire, y devuelve el número de la siguiente: ahí
    va la primera tabla. Con dos renglones, como en los listados, es la 5.
    """
    hoja.cell(row=1, column=1, value=titulo).font = Font(bold=True, size=14, color='8C5A66')
    for numero, renglon in enumerate(renglones, start=2):
        hoja.cell(row=numero, column=1, value=renglon)
    return len(renglones) + 3


def escribir_tabla_excel(hoja, fila_titulos, columnas, filas):
    """Una tabla dentro de la hoja: los títulos de las columnas en la fila
    fila_titulos, en negrita sobre rosa claro, y las filas debajo.

    Los importes y las fechas van como números y fechas, para que se puedan
    sumar y ordenar en Excel. Los textos van siempre como texto. Devuelve el
    número de la primera fila libre después de la tabla.
    """
    for posicion, columna in enumerate(columnas, start=1):
        celda = hoja.cell(row=fila_titulos, column=posicion, value=columna.titulo)
        celda.font = Font(bold=True, color='8C5A66')
        celda.fill = PatternFill('solid', fgColor='F0E2E4')

    for numero_de_fila, fila in enumerate(filas, start=fila_titulos + 1):
        for posicion, (valor, columna) in enumerate(zip(fila, columnas), start=1):
            celda = hoja.cell(row=numero_de_fila, column=posicion)
            if columna.tipo == 'texto':
                # data_type 's' lo deja como texto aunque empiece con =, +,
                # - o @: si no, Excel lo tomaría como una fórmula y un nombre
                # cargado como "=..." podría ejecutar algo al abrir la planilla.
                celda.value = '' if valor is None else str(valor)
                celda.data_type = 's'
            else:
                celda.value = valor
                celda.number_format = FORMATO[columna.tipo]
            celda.alignment = Alignment(vertical='top', wrap_text=columna.tipo == 'texto')

    return fila_titulos + 1 + len(filas)


def ajustar_anchos_excel(hoja, columnas):
    """El ancho de cada columna de la hoja, sacado del ancho de la Columna."""
    for posicion, columna in enumerate(columnas, start=1):
        hoja.column_dimensions[get_column_letter(posicion)].width = columna.ancho * CARACTERES_POR_ANCHO


def guardar_excel(libro):
    """El libro en memoria, listo para la respuesta."""
    archivo = BytesIO()
    libro.save(archivo)
    archivo.seek(0)
    return archivo


def armar_excel(titulo, fecha, cantidad, texto_filtros, columnas, filas):
    """La planilla .xlsx de un listado: título, detalle, encabezado y filas."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo[:31]  # Excel no acepta nombres de hoja más largos.

    fila_titulos = escribir_encabezado_excel(hoja, titulo, [f'{fecha} · {cantidad}', texto_filtros])
    escribir_tabla_excel(hoja, fila_titulos, columnas, filas)
    ajustar_anchos_excel(hoja, columnas)

    # Todo lo que está arriba de la primera fila de datos (el título, los
    # filtros y los títulos de las columnas) queda fijo al desplazar la hoja.
    hoja.freeze_panes = hoja.cell(row=fila_titulos + 1, column=1)

    return guardar_excel(libro)


# ---------------------------------------------------------------------
# La acción que heredan los ViewSets
# ---------------------------------------------------------------------

class ExportarMixin:
    """Suma a un ViewSet la acción exportar (corrección E10 del tutor).

    El ViewSet declara:
      EXPORTAR_TITULO        el título del archivo ('Pedidos')
      EXPORTAR_ARCHIVO       el principio del nombre del archivo ('pedidos')
      EXPORTAR_UNIDAD        el singular y el plural ('pedido', 'pedidos')
      COLUMNAS_PARA_EXPORTAR la lista de Columna
      describir_filtros()    los filtros de la URL en palabras
      registros_para_exportar()  opcional, si cambia el orden de la pantalla

    Y su get_queryset() tiene que filtrar también cuando la acción es
    'exportar', no solo 'list'.
    """

    def describir_filtros(self, parametros):
        """Los filtros aplicados en palabras. Sin sobrescribir, ninguno."""
        return []

    def registros_para_exportar(self):
        """Los registros del archivo: lo mismo que hace el list() de DRF,
        sin la paginación. Un ViewSet lo sobrescribe si la pantalla los
        muestra en otro orden, como Materiales con sus discontinuados."""
        return list(self.filter_queryset(self.get_queryset()))

    #   GET /api/<listado>/exportar/?formato=pdf&<los filtros del listado>

    @action(detail=False, methods=['get'])
    def exportar(self, request):
        """Devuelve el listado filtrado como PDF o como Excel (corrección E10)."""
        formato = request.query_params.get('formato')
        if formato not in ('pdf', 'excel'):
            raise ParseError('El formato tiene que ser pdf o excel.')

        registros = self.registros_para_exportar()
        columnas = self.COLUMNAS_PARA_EXPORTAR
        filas = [[columna.valor(registro) for columna in columnas] for registro in registros]

        hoy = timezone.localdate()
        singular, plural = self.EXPORTAR_UNIDAD
        cantidad = f'{len(filas)} {singular if len(filas) == 1 else plural}'
        filtros = self.describir_filtros(request.query_params)
        texto_filtros = ' · '.join(filtros) if filtros else 'Sin filtros'
        datos = (self.EXPORTAR_TITULO, escribir_fecha(hoy), cantidad, texto_filtros, columnas, filas)

        nombre = f'{self.EXPORTAR_ARCHIVO}-{hoy.isoformat()}'

        # El PDF se muestra en la pestaña (inline), como el comprobante. El
        # Excel se descarga (attachment): el navegador no lo puede mostrar.
        if formato == 'pdf':
            return FileResponse(armar_pdf(*datos), as_attachment=False, filename=f'{nombre}.pdf')

        return FileResponse(armar_excel(*datos), as_attachment=True, filename=f'{nombre}.xlsx')
